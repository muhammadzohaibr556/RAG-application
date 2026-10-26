"""Generate small real documents to verify extraction and original citation locations."""

from io import BytesIO

import pymupdf
import pytest
from docx import Document

from app.errors import InvalidUpload, UploadTooLarge
from app.extraction import extract, extract_docx, extract_pdf, extract_txt
from app.normalization import normalize


def test_pdf_pages():
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Page one: permit")
        pdf.new_page().insert_text((72, 72), "Page two: inspection")
        units = extract_pdf(pdf.tobytes(), 10000)
    assert [unit.position for unit in units] == [1, 2]
    assert "inspection" in units[1].text


def test_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("")
    doc.add_paragraph("Third paragraph")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table text"
    file = BytesIO()
    doc.save(file)
    units = normalize(extract_docx(file.getvalue(), 10000))
    assert [unit.position for unit in units] == [1, 3, 4]
    assert units[-1].text == "Table text"


def test_txt_line_positions_and_unicode():
    units = normalize(extract_txt(b"\xef\xbb\xbfFirst\r\n\r\nThird\t line", 1000))
    assert [(unit.position, unit.text) for unit in units] == [(1, "First"), (3, "Third line")]


@pytest.mark.parametrize(
    "content,extension",
    [
        (b"broken", ".pdf"),
        (b"not a zip", ".docx"),
        (b"\xff", ".txt"),
        (b"\x00", ".txt"),
        (b"  \n", ".txt"),
    ],
)
def test_invalid_content(content, extension):
    with pytest.raises(InvalidUpload):
        extract(content, extension, 10000)


def test_scanned_pdf_is_rejected():
    with pymupdf.open() as pdf:
        pdf.new_page()
        with pytest.raises(InvalidUpload, match="No readable text"):
            extract(pdf.tobytes(), ".pdf", 1000)


def test_text_limit():
    with pytest.raises(UploadTooLarge):
        extract_txt(b"too long", 3)
